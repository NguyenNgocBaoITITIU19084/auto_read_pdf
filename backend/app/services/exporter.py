import io
import re

# Column colours offered by export templates: strong header shade, light body tint, readable header text
EXPORT_PALETTE = {
    "red": ("C0392B", "FDECEA", "FFFFFF"),
    "orange": ("E67E22", "FDF2E6", "FFFFFF"),
    "yellow": ("F1C40F", "FEF9E3", "1F2937"),
    "green": ("27AE60", "E8F6EE", "FFFFFF"),
    "teal": ("16A085", "E6F5F2", "FFFFFF"),
    "blue": ("2980B9", "E8F1F9", "FFFFFF"),
    "purple": ("8E44AD", "F3EAF7", "FFFFFF"),
    "pink": ("D63384", "FBE9F2", "FFFFFF"),
    "gray": ("7F8C8D", "F1F3F4", "FFFFFF"),
}
EXPORT_FILLS = ("header", "column")
_HEX = re.compile(r"^#?([0-9A-Fa-f]{6})$")


def normalize_export_color(value) -> str | None:
    """A palette id, or a custom colour as '#RRGGBB'; anything else is no colour."""
    if value in EXPORT_PALETTE:
        return value
    m = _HEX.match(str(value or "").strip())
    return f"#{m.group(1).upper()}" if m else None


def color_shades(value) -> tuple[str, str, str] | None:
    """(header, body tint, header text) for a palette id or a '#RRGGBB' colour. A custom colour's tint is it mixed
    85% with white, and its header text is dark on light colours, white otherwise."""
    if value in EXPORT_PALETTE:
        return EXPORT_PALETTE[value]
    color = normalize_export_color(value)
    if not color:
        return None
    r, g, b = (int(color[i:i + 2], 16) for i in (1, 3, 5))
    tint = "".join(f"{round(c + (255 - c) * 0.85):02X}" for c in (r, g, b))
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return color[1:], tint, "1F2937" if luminance > 0.6 else "FFFFFF"

_DEFAULT_HEADER = "2C3E50"
_ROW_TINT = "F2F7FA"

# Short codes and numbers read better centred (decided by column key, so a renamed header keeps it)
_CENTER_KEYS = {
    "STT", "Q'ty", "Booking No", "Block", "Equipment Type", "ETD", "Trạng thái",
    "site_id", "in_out_voyage", "fel", "iso", "event_type", "vgm", "cust", "haz", "temp",
}
_CENTER_LABELS = {"STT", "No.", "Q'ty", "Booking No", "Số Booking", "Block", "Equipment Type", "Loại cont", "ETD",
                  "Ngày tàu chạy"}


def export_columns_to_excel(rows: list[dict], columns: list[dict]) -> io.BytesIO:
    """Excel file of `rows` (dicts keyed by column key) with `columns` in order, each
    {"key", "header", "color": palette id | "#RRGGBB" | None, "fill": "header" | "column"}.

    "header" colours the header cell only; "column" also tints every body cell of that column (instead of the
    alternating stripe). Blank / None values are written as "null", like the table shows them."""
    if not rows or not columns:
        raise ValueError("Data and columns cannot be empty")

    # Heavy imports are deferred so backend startup stays fast
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    def fill(hex_color: str) -> PatternFill:
        return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

    side = Side(style="thin", color="CCCCCC")
    border = Border(left=side, right=side, top=side, bottom=side)
    cell_font = Font(name="Segoe UI", size=10)

    wb = Workbook()
    ws = wb.active
    ws.row_dimensions[1].height = 28
    widths = []
    for c_idx, col in enumerate(columns, start=1):
        palette = color_shades(col.get("color"))
        header_hex, tint_hex, font_hex = palette if palette else (_DEFAULT_HEADER, None, "FFFFFF")
        header = str(col.get("header") or col["key"])
        cell = ws.cell(row=1, column=c_idx, value=header)
        cell.fill = fill(header_hex)
        cell.font = Font(name="Segoe UI", size=11, bold=True, color=font_hex)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

        body_fill = fill(tint_hex) if palette and col.get("fill") == "column" else None
        center = col["key"] in _CENTER_KEYS
        width = len(header)
        for r_idx, row in enumerate(rows, start=2):
            value = row.get(col["key"])
            if value is None or value == "":
                value = "null"
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = cell_font
            cell.border = border
            cell.fill = body_fill or fill(_ROW_TINT if c_idx % 2 == 0 else "FFFFFF")
            cell.alignment = Alignment(horizontal="center" if center else "left", vertical="center")
            width = max(width, *(len(line) for line in str(value).split("\n")))
        widths.append(width)

    for r_idx in range(2, len(rows) + 2):
        ws.row_dimensions[r_idx].height = 20
    for c_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=c_idx).column_letter].width = max(width + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_to_excel_buffer(data: list[dict], selected_columns: list[str]) -> io.BytesIO:
    """Older call: rows keyed by header text, plain colours."""
    if not data or not selected_columns:
        raise ValueError("Data and selected columns cannot be empty")
    columns = [{"key": name, "header": name, "color": None, "fill": "header"} for name in selected_columns]
    out = export_columns_to_excel(data, columns)
    # keep the old label-based centring for callers that still send labels
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    wb = load_workbook(out)
    ws = wb.active
    for c_idx, name in enumerate(selected_columns, start=1):
        if name in _CENTER_LABELS:
            for r_idx in range(2, ws.max_row + 1):
                ws.cell(row=r_idx, column=c_idx).alignment = Alignment(horizontal="center", vertical="center")
    result = io.BytesIO()
    wb.save(result)
    result.seek(0)
    return result
