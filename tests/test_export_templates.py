"""Excel export templates: storage, validation and the styled export they drive (read back with openpyxl)."""
import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.core import database as db
from backend.app.main import app
from backend.app.services.exporter import EXPORT_PALETTE, export_columns_to_excel


@pytest.fixture(autouse=True)
def _db():
    db.init_db()


def uniq(name):
    return f"{name} {time.time_ns()}"


ROWS = [
    {"STT": 1, "Booking No": "BK1", "Carrier": "ONE", "Vessel": "EVER OATH", "Ghi chú": ""},
    {"STT": 2, "Booking No": "BK2", "Carrier": "MSC", "Vessel": None, "Ghi chú": "gấp"},
]


def _sheet(buf):
    return load_workbook(io.BytesIO(buf.getvalue())).active


def _rgb(cell):
    return cell.fill.start_color.rgb[-6:]


# --- styled export -------------------------------------------------------------------------------------------
def test_export_follows_column_order_headers_and_colours():
    ws = _sheet(export_columns_to_excel(ROWS, [
        {"key": "Vessel", "header": "Tàu", "color": "blue", "fill": "column"},
        {"key": "Booking No", "header": "Số BK", "color": "red", "fill": "header"},
        {"key": "STT", "header": "", "color": None, "fill": "header"},
    ]))
    assert [c.value for c in ws[1]] == ["Tàu", "Số BK", "STT"]
    assert [c.value for c in ws[2]] == ["EVER OATH", "BK1", 1]
    assert ws["A3"].value == "null"                                  # blank -> "null"
    blue_header, blue_tint, _ = EXPORT_PALETTE["blue"]
    red_header, red_tint, _ = EXPORT_PALETTE["red"]
    assert _rgb(ws["A1"]) == blue_header and _rgb(ws["A2"]) == blue_tint and _rgb(ws["A3"]) == blue_tint
    assert _rgb(ws["B1"]) == red_header and _rgb(ws["B2"]) != red_tint   # header only
    assert _rgb(ws["C1"]) == "2C3E50"                                # default header
    assert ws["B2"].alignment.horizontal == "center"                 # by key, whatever the header says


def test_readable_header_text_on_yellow():
    ws = _sheet(export_columns_to_excel(ROWS, [{"key": "Carrier", "header": "Hãng", "color": "yellow", "fill": "header"}]))
    assert ws["A1"].font.color.rgb[-6:] == EXPORT_PALETTE["yellow"][2]


def test_two_columns_renamed_alike_keep_their_own_values():
    ws = _sheet(export_columns_to_excel(ROWS, [
        {"key": "Booking No", "header": "Mã", "color": None, "fill": "header"},
        {"key": "Carrier", "header": "Mã", "color": None, "fill": "header"},
    ]))
    assert [c.value for c in ws[2]] == ["BK1", "ONE"]


def test_export_endpoint_keyed_columns_and_filename():
    client = TestClient(app)
    res = client.post("/api/v1/export/excel", json={
        "data": ROWS, "filename": "Kế hoạch 16-10.xlsx",
        "columns": [{"key": "Booking No", "header": "Số BK", "color": "green", "fill": "column"}],
    })
    assert res.status_code == 200
    assert "filename*=UTF-8''K%E1%BA%BF" in res.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(res.content)).active
    assert ws["A1"].value == "Số BK" and ws["A3"].value == "BK2"
    old = client.post("/api/v1/export/excel", json={"data": [{"Số Booking": "X"}], "selected_columns": ["Số Booking"]})
    assert old.status_code == 200


# --- templates -----------------------------------------------------------------------------------------------
COLS = [{"key": "Booking No", "header": "Số BK", "color": "red", "fill": "column"},
        {"key": "Vessel", "header": "", "color": None, "fill": "everything"},
        {"key": "Booking No", "header": "dup"}]


def test_template_crud_and_cleaning():
    name = uniq("Kế hoạch đội xe")
    t = db.create_export_template("booking", name, COLS, "ke_hoach_{date}")
    assert t["columns"] == [{"key": "Booking No", "header": "Số BK", "color": "red", "fill": "column"},
                            {"key": "Vessel", "header": "", "color": None, "fill": "header"}]
    assert t["filename"] == "ke_hoach_{date}"
    assert t["id"] in [x["id"] for x in db.list_export_templates("booking")]
    assert t["id"] not in [x["id"] for x in db.list_export_templates("vessel")]

    u = db.update_export_template(t["id"], name + " v2", [{"key": "Carrier", "header": "Hãng", "color": "blue", "fill": "header"}], "")
    assert u["name"] == name + " v2" and [c["key"] for c in u["columns"]] == ["Carrier"]
    assert db.update_export_template(999999999, "x", COLS) is None
    assert db.delete_export_template(t["id"]) is True
    assert db.delete_export_template(t["id"]) is False


def test_template_validation():
    name = uniq("Mẫu A")
    db.create_export_template("vessel", name, COLS)
    with pytest.raises(ValueError, match="Đã có mẫu"):
        db.create_export_template("vessel", name.upper(), COLS)          # case-insensitive per tab
    db.create_export_template("container", name, COLS)                   # other tab is fine
    with pytest.raises(ValueError):
        db.create_export_template("vessel", "   ", COLS)
    with pytest.raises(ValueError):
        db.create_export_template("vessel", uniq("B"), [])
    with pytest.raises(ValueError):
        db.create_export_template("nowhere", uniq("C"), COLS)


def test_template_api():
    client = TestClient(app)
    name = uniq("API mẫu")
    created = client.post("/api/v1/export/templates", json={"tab": "container", "name": name, "columns": COLS})
    assert created.status_code == 200
    tid = created.json()["id"]
    assert client.post("/api/v1/export/templates", json={"tab": "container", "name": name, "columns": COLS}).status_code == 400
    listed = client.get("/api/v1/export/templates", params={"tab": "container"}).json()
    assert tid in [x["id"] for x in listed]
    renamed = client.put(f"/api/v1/export/templates/{tid}", json={"name": name + "!", "columns": COLS, "filename": "f"})
    assert renamed.status_code == 200 and renamed.json()["filename"] == "f"
    assert client.delete(f"/api/v1/export/templates/{tid}").status_code == 200
    assert client.delete(f"/api/v1/export/templates/{tid}").status_code == 404
    assert client.get("/api/v1/export/templates", params={"tab": "bad"}).status_code == 400
    palette = client.get("/api/v1/export/palette").json()
    assert {"id": "red", "header": "#C0392B", "tint": "#FDECEA", "text": "#FFFFFF"} in palette


def test_custom_hex_colours():
    from backend.app.services.exporter import color_shades
    assert color_shades("#1e88e5") == ("1E88E5", "DDEDFB", "FFFFFF")
    assert color_shades("FFEB3B")[2] == "1F2937"                       # light colour -> dark header text
    assert color_shades("#12345") is None and color_shades("nope") is None
    ws = _sheet(export_columns_to_excel(ROWS, [{"key": "Carrier", "header": "Hãng", "color": "#1E88E5", "fill": "column"}]))
    assert _rgb(ws["A1"]) == "1E88E5" and _rgb(ws["A2"]) == "DDEDFB"
    t = db.create_export_template("booking", uniq("Hex"), [{"key": "Carrier", "color": "1e88e5", "fill": "column"}])
    assert t["columns"][0]["color"] == "#1E88E5"


def test_unreadable_colour_is_refused_not_dropped():
    with pytest.raises(ValueError, match="Màu không hợp lệ"):
        db.create_export_template("booking", uniq("Bad colour"), [{"key": "STT", "color": "#00897", "fill": "column"}])
    client = TestClient(app)
    res = client.post("/api/v1/export/templates", json={"tab": "booking", "name": uniq("Bad"), "columns": [{"key": "STT", "color": "teal-ish"}]})
    assert res.status_code == 400 and "Màu không hợp lệ" in res.json()["detail"]
    ok = db.create_export_template("booking", uniq("No colour"), [{"key": "STT", "color": None}, {"key": "Carrier", "color": ""}])
    assert [c["color"] for c in ok["columns"]] == [None, None]
